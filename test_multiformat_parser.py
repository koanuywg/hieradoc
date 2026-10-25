import os
import sys
from dotenv import load_dotenv

load_dotenv("D:/code/HỌC/hieradoc/.env")
from hieradoc import HieraDoc
from hieradoc.core.extractor import FileExtractor

def test_universal_formats():
    print("=====================================================================")
    print(">>> TEST: MULTI-FORMAT PARSER (DOCX via MarkItDown & Scanned Image via Vision OCR)")
    print("=====================================================================")

    # 1. Tạo file Word (.docx) mẫu
    import docx
    doc_path = "sample_test_contract.docx"
    doc = docx.Document()
    doc.add_heading("HỢP ĐỒNG CUNG ỨNG DỊCH VỤ ĐIỆN TOÁN ĐÁM MÂY", 0)
    doc.add_heading("Chương I: ĐIỀU KHOẢN CHUNG VÀ CAM KẾT", level=1)
    doc.add_paragraph("Điều 1. Phạm vi cung cấp dịch vụ")
    doc.add_paragraph("Bên A đồng ý cung cấp giải pháp lưu trữ dữ liệu đám mây cho Bên B với dung lượng 50TB.")
    doc.add_heading("Chương II: PHÍ VÀ THANH TOÁN", level=1)
    doc.add_paragraph("Điều 2. Thời hạn thanh toán định kỳ")
    doc.add_paragraph("Bên B phải hoàn tất chuyển khoản tiền dịch vụ vào ngày 05 hàng tháng với số tiền 2000 USD.")
    doc.save(doc_path)
    print(f"Created sample docx: {doc_path}")

    # 2. Tạo ảnh mẫu giả lập văn bản scan (.png)
    from PIL import Image, ImageDraw, ImageFont
    img_path = "sample_scanned_receipt.png"
    img = Image.new("RGB", (800, 300), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    # Ghi text vào ảnh
    draw.text((30, 30), "BIÊN BẢN NGHIỆM THU DỰ ÁN CÔNG NGHỆ THÔNG TIN", fill=(0, 0, 0))
    draw.text((30, 80), "Chương I: KẾT QUẢ NGHIỆM THU", fill=(0, 0, 0))
    draw.text((30, 120), "Điều 1: Đánh giá chất lượng hệ thống phần mềm", fill=(0, 0, 0))
    draw.text((30, 160), "Hệ thống đạt chuẩn an toàn thông tin ISO 27001 và vận hành ổn định 99.9%.", fill=(0, 0, 0))
    draw.text((30, 210), "Chi phí quyết toán cuối cùng là 45,000 USD.", fill=(0, 0, 0))
    img.save(img_path)
    print(f"Created sample scanned image: {img_path}")

    db_path = "test_multiformat.db"
    if os.path.exists(db_path):
        os.remove(db_path)

    hieradoc = HieraDoc(storage="sqlite", db_path=db_path)

    # 3. Test Ingesting DOCX (MarkItDown)
    print("\n--- 1. Testing Word (.docx) Ingestion via MarkItDown ---")
    docx_res = hieradoc.index(file_path_or_text=doc_path, doc_id="DOCX_CONTRACT", cluster="Hợp Đồng Mẫu")
    print("DOCX Index Result:", docx_res)

    # 4. Test Ingesting Image (VisionOCRAgent)
    print("\n--- 2. Testing Scanned Image (.png) Ingestion via VisionOCRAgent ---")
    img_res = hieradoc.index(file_path_or_text=img_path, doc_id="SCANNED_RECEIPT", cluster="Biên Bản Nghiệm Thu")
    print("Image Index Result:", img_res)

    # 5. Query Word Document
    print("\n--- 3. Querying Word (.docx) Document ---")
    q1 = "Hạn thanh toán tiền dịch vụ hàng tháng là ngày nào và số tiền bao nhiêu?"
    r1 = hieradoc.query(question=q1)
    print("Answer 1:\n", r1.answer)
    print("Citations 1:")
    for c in r1.citations:
        print(f" - [{c.doc_id}] Lines {c.lines}: \"{c.exact_quote}\"")

    assert "ngày 05" in r1.answer or "05" in r1.answer or "2000" in r1.answer, "Failed to retrieve payment terms from DOCX!"

    # 6. Query Scanned Image Document
    print("\n--- 4. Querying Scanned Image (.png) Document ---")
    q2 = "Hệ thống đạt chuẩn an toàn thông tin gì và chi phí quyết toán là bao nhiêu?"
    r2 = hieradoc.query(question=q2)
    print("Answer 2:\n", r2.answer)
    print("Citations 2:")
    for c in r2.citations:
        print(f" - [{c.doc_id}] Lines {c.lines}: \"{c.exact_quote}\"")

    assert "27001" in r2.answer or "45,000" in r2.answer or "45000" in r2.answer, "Failed to retrieve ISO or settlement from Scanned PNG!"

    print("\n>>> ALL TESTS PASSED! Multi-format support (MarkItDown + Vision OCR) successfully validated!")

if __name__ == "__main__":
    test_universal_formats()
